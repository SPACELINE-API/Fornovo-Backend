import re
import io
from dataclasses import dataclass, field
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# paleta
class Cor:
    VERDE_ESCURO  = "1B4332"
    VERDE_MEDIO   = "2D6A4F"
    VERDE_CLARO   = "D8F3DC"
    VERDE_TEXTO   = "1B4332"
    VERMELHO_BG   = "FFE5E5"
    VERMELHO_TXT  = "7B1D1D"
    AMARELO_BG    = "FFF9C4"
    AMARELO_TXT   = "5C4A00"
    CINZA_CLARO   = "F4F4F4"
    CINZA_BORDA   = "CCCCCC"
    BRANCO        = "FFFFFF"
    TEXTO_NORMAL  = "1A1A1A"
    TEXTO_SUAVE   = "555555"


def hex_to_rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)


def set_cell_border(cell, *, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        el = OxmlElement(f"w:{side}")
        if val is None:
            el.set(qn("w:val"), "none")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(val.get("sz", 6)))
            el.set(qn("w:color"), val.get("color", "000000").lstrip("#"))
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tcPr.append(m)


def set_run_color(run, hex_color):
    r, g, b = hex_to_rgb(hex_color)
    run.font.color.rgb = RGBColor(r, g, b)


def add_horizontal_rule(doc, cor=Cor.VERDE_ESCURO, espaco_antes=6, espaco_depois=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(espaco_antes)
    p.paragraph_format.space_after = Pt(espaco_depois)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), cor.lstrip("#"))
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_spacer(doc, pt_before=8, pt_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(pt_before)
    p.paragraph_format.space_after = Pt(pt_after)
    return p

@dataclass
class Item:
    emoji: str = ""
    codigo: str = ""
    descricao: str = ""
    status: str = ""
    justificativa: str = ""
    recomendacao: str = ""
    layers_confirmados: str = ""
    layers_ausentes: str = ""
    evidencias: list = field(default_factory=list)


@dataclass
class Grupo:
    titulo: str = ""
    resumo: str = ""
    itens: list = field(default_factory=list)


@dataclass
class Sumario:
    conforme: int = 0
    nao_conforme: int = 0
    inconclusivo: int = 0
    perc_conf: str = "0%"
    perc_nconf: str = "0%"
    perc_incon: str = "0%"
    total: int = 0


def parse_relatorio(md: str):
    lines = md.split("\n")
    meta = {}
    grupos = []
    grupo_atual = None
    item_atual = None
    secao = None
    sumario = Sumario()

    for line in lines:
        if line.startswith("**Arquivo analisado:**"):
            m = re.search(r"`(.+?)`", line)
            meta["arquivo"] = m.group(1) if m else line.split(":**")[1].strip()
        if line.startswith("**Data/Hora:**"):
            meta["data_hora"] = line.split(":**")[1].strip()
        if line.startswith("**Normas de referência:**"):
            meta["normas"] = line.split(":**")[1].strip()

        m = re.search(r"✅ Conforme\s+\|\s+(\d+)\s+\|\s+([\d%]+)", line)
        if m:
            sumario.conforme = int(m.group(1)); sumario.perc_conf = m.group(2)
        m = re.search(r"❌ Não Conforme\s+\|\s+(\d+)\s+\|\s+([\d%]+)", line)
        if m:
            sumario.nao_conforme = int(m.group(1)); sumario.perc_nconf = m.group(2)
        m = re.search(r"⚠️\s+Inconclusivo\s+\|\s+(\d+)\s+\|\s+([\d%]+)", line)
        if m:
            sumario.inconclusivo = int(m.group(1)); sumario.perc_incon = m.group(2)

        m = re.match(r"^## (.+?)\s+\*\((.+?)\)\*", line)
        if m:
            grupo_atual = Grupo(titulo=m.group(1), resumo=m.group(2))
            grupos.append(grupo_atual)
            item_atual = None; secao = None
            continue

        m = re.match(r"^### (✅|❌|⚠️)\s+\[(.+?)\]\s+(.+)", line)
        if m and grupo_atual:
            item_atual = Item(
                emoji=m.group(1), codigo=m.group(2),
                descricao=re.sub(r"\*\*", "", m.group(3))
            )
            grupo_atual.itens.append(item_atual)
            secao = None
            continue

        if not item_atual:
            continue

        if line.startswith("**Status:**"):
            item_atual.status = line.split(":**")[1].strip().replace("`", "")
        elif line.startswith("**Justificativa:**"):
            item_atual.justificativa = line.split(":**", 1)[1].strip()
        elif line.startswith("**Recomendação:**"):
            item_atual.recomendacao = line.split(":**", 1)[1].strip()
        elif line.startswith("**Layers confirmados:**"):
            item_atual.layers_confirmados = line.split(":**", 1)[1].strip().replace("`", "")
        elif line.startswith("**Layers não encontrados:**"):
            item_atual.layers_ausentes = line.split(":**", 1)[1].strip().replace("`", "")
        elif line.startswith("**Evidências"):
            secao = "evidencias"
        elif secao == "evidencias" and line.startswith("- "):
            item_atual.evidencias.append(line[2:].strip())

    sumario.total = sumario.conforme + sumario.nao_conforme + sumario.inconclusivo
    return meta, sumario, grupos


# status
def status_info(emoji):
    if emoji == "✅":
        return dict(label="Conforme",      bg=Cor.VERDE_CLARO,  txt=Cor.VERDE_TEXTO,  borda=Cor.VERDE_MEDIO)
    if emoji == "❌":
        return dict(label="Não Conforme",  bg=Cor.VERMELHO_BG,  txt=Cor.VERMELHO_TXT, borda="C0392B")
    return         dict(label="Inconclusivo", bg=Cor.AMARELO_BG, txt=Cor.AMARELO_TXT, borda="E67E22")


# badge
def adicionar_badge(doc, texto, bg, txt_cor):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    set_cell_bg(cell, bg)
    set_cell_border(cell)
    set_cell_margins(cell, top=50, bottom=50, left=140, right=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(10)
    set_run_color(run, txt_cor)
    return t


# metadados
def adicionar_meta(doc, meta):
    dados = [
        ("📄 Arquivo",   meta.get("arquivo", "—")),
        ("📅 Data/Hora", meta.get("data_hora", "—")),
        ("📋 Normas",    meta.get("normas", "—")),
    ]
    t = doc.add_table(rows=len(dados), cols=2)
    t.style = "Table Grid"
    larguras = [Inches(1.8), Inches(5.5)]
    for i, (label, valor) in enumerate(dados):
        row = t.rows[i]
        # coluna label
        c0 = row.cells[0]
        set_cell_bg(c0, Cor.CINZA_CLARO)
        set_cell_margins(c0, top=70, bottom=70, left=160, right=160)
        p = c0.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        set_run_color(run, Cor.VERDE_ESCURO)
        # coluna valor
        c1 = row.cells[1]
        set_cell_margins(c1, top=70, bottom=70, left=160, right=160)
        p = c1.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(valor)
        run.font.size = Pt(10)
        set_run_color(run, Cor.TEXTO_NORMAL)
    return t


# tabela
def adicionar_sumario(doc, sumario):
    linhas = [
        ("Status",        "Qtd.",                        "%",                   Cor.VERDE_ESCURO, Cor.BRANCO,       True),
        ("✅  Conforme",  str(sumario.conforme),          sumario.perc_conf,     Cor.VERDE_CLARO,  Cor.VERDE_TEXTO,  True),
        ("❌  Não Conforme", str(sumario.nao_conforme),  sumario.perc_nconf,    Cor.VERMELHO_BG,  Cor.VERMELHO_TXT, True),
        ("⚠️  Inconclusivo", str(sumario.inconclusivo),  sumario.perc_incon,    Cor.AMARELO_BG,   Cor.AMARELO_TXT,  True),
        ("Total",         str(sumario.total),             "100%",                Cor.CINZA_CLARO,  Cor.TEXTO_NORMAL, True),
    ]
    t = doc.add_table(rows=len(linhas), cols=3)
    t.style = "Table Grid"
    for i, (col1, col2, col3, bg, txt, bold) in enumerate(linhas):
        row = t.rows[i]
        for j, texto in enumerate([col1, col2, col3]):
            cell = row.cells[j]
            set_cell_bg(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(texto)
            run.bold = bold
            run.font.size = Pt(11 if i == 0 else 10)
            set_run_color(run, txt)
    return t


# cabeçalho
def adicionar_cabecalho_grupo(doc, grupo):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    set_cell_bg(cell, Cor.VERDE_ESCURO)
    set_cell_border(cell, top={"sz": 0}, bottom={"sz": 0}, left={"sz": 0}, right={"sz": 0})
    set_cell_margins(cell, top=140, bottom=140, left=240, right=240)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(grupo.titulo.upper())
    run.bold = True
    run.font.size = Pt(13)
    set_run_color(run, Cor.BRANCO)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    run2 = p2.add_run(grupo.resumo)
    run2.italic = True
    run2.font.size = Pt(10)
    set_run_color(run2, "A8D5B5")
    return t


# itens
def adicionar_card_item(doc, item, is_zebra=False):
    info = status_info(item.emoji)
    bg = Cor.CINZA_CLARO if is_zebra else Cor.BRANCO
    borda_cor = info["borda"]

    t = doc.add_table(rows=2, cols=1)
    t.style = "Table Grid"

    c_header = t.cell(0, 0)
    set_cell_bg(c_header, bg)
    set_cell_border(c_header,
        top={"sz": 12, "color": borda_cor},
        bottom={"sz": 0},
        left={"sz": 24, "color": borda_cor},
        right={"sz": 0},
    )
    set_cell_margins(c_header, top=120, bottom=60, left=200, right=160)
    ph = c_header.paragraphs[0]
    ph.paragraph_format.space_before = Pt(0)
    ph.paragraph_format.space_after = Pt(0)
    r1 = ph.add_run(f"[{item.codigo}]  ")
    r1.bold = True; r1.font.size = Pt(11)
    set_run_color(r1, Cor.VERDE_MEDIO)
    r2 = ph.add_run(item.descricao)
    r2.bold = True; r2.font.size = Pt(11)
    set_run_color(r2, Cor.TEXTO_NORMAL)

    c_body = t.cell(1, 0)
    set_cell_bg(c_body, bg)
    set_cell_border(c_body,
        top={"sz": 0},
        bottom={"sz": 6, "color": borda_cor},
        left={"sz": 24, "color": borda_cor},
        right={"sz": 0},
    )
    set_cell_margins(c_body, top=60, bottom=140, left=200, right=160)

    pb = c_body.paragraphs[0]
    pb.paragraph_format.space_before = Pt(0)
    pb.paragraph_format.space_after = Pt(0)
    rb = pb.add_run(f"  {item.emoji}  {info['label']}  ")
    rb.bold = True; rb.font.size = Pt(9)
    set_run_color(rb, info["txt"])
    rPr = rb._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), info["bg"].lstrip("#"))
    rPr.append(shd)

    def campo(label, valor):
        p = c_body.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(3)
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.bold = True; r_lbl.font.size = Pt(10)
        set_run_color(r_lbl, Cor.VERDE_ESCURO)
        r_val = p.add_run(valor)
        r_val.font.size = Pt(10)
        set_run_color(r_val, Cor.TEXTO_NORMAL)

    if item.justificativa:
        campo("Justificativa", item.justificativa)
    if item.recomendacao and item.recomendacao.lower() not in ("nenhuma", "nenhum"):
        campo("Recomendação", item.recomendacao)
    if item.layers_confirmados:
        campo("Layers confirmados", item.layers_confirmados)
    if item.layers_ausentes:
        campo("Layers ausentes", item.layers_ausentes)
    if item.evidencias:
        p_ev = c_body.add_paragraph()
        p_ev.paragraph_format.space_before = Pt(6)
        p_ev.paragraph_format.space_after = Pt(2)
        r_ev = p_ev.add_run("Evidências:")
        r_ev.bold = True; r_ev.font.size = Pt(10)
        set_run_color(r_ev, Cor.VERDE_ESCURO)
        for ev in item.evidencias:
            p_i = c_body.add_paragraph()
            p_i.paragraph_format.space_before = Pt(2)
            p_i.paragraph_format.space_after = Pt(2)
            p_i.paragraph_format.left_indent = Inches(0.25)
            r_bul = p_i.add_run("▸ ")
            r_bul.font.size = Pt(10)
            set_run_color(r_bul, Cor.VERDE_MEDIO)
            r_txt = p_i.add_run(ev)
            r_txt.font.size = Pt(10)
            set_run_color(r_txt, Cor.TEXTO_NORMAL)
    return t


# titulo
def adicionar_titulo_secao(doc, texto, pt_size=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "8")
    bot.set(qn("w:color"), Cor.VERDE_ESCURO.lstrip("#"))
    pBdr.append(bot); pPr.append(pBdr)
    run = p.add_run(texto)
    run.bold = True; run.font.size = Pt(pt_size)
    set_run_color(run, Cor.VERDE_ESCURO)
    return p


def gerar_docx_bytes(markdown_text: str) -> bytes:
    meta, sumario, grupos = parse_relatorio(markdown_text)
    doc = Document()

    for sec in doc.sections:
        sec.top_margin    = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin   = Inches(1)
        sec.right_margin  = Inches(1)

    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)

    p_titulo = doc.add_paragraph()
    p_titulo.paragraph_format.space_before = Pt(0)
    p_titulo.paragraph_format.space_after = Pt(4)
    r = p_titulo.add_run("RELATÓRIO DE CONFORMIDADE")
    r.bold = True; r.font.size = Pt(22)
    set_run_color(r, Cor.VERDE_ESCURO)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Análise técnica conforme normas NBR")
    r_sub.italic = True; r_sub.font.size = Pt(12)
    set_run_color(r_sub, Cor.TEXTO_SUAVE)
    add_horizontal_rule(doc, Cor.VERDE_ESCURO, espaco_antes=0, espaco_depois=14)
    adicionar_titulo_secao(doc, "Informações do Documento", pt_size=13)
    adicionar_meta(doc, meta)
    add_spacer(doc, 12, 6)
    adicionar_titulo_secao(doc, "Sumário Executivo", pt_size=13)
    adicionar_sumario(doc, sumario)
    add_spacer(doc, 16, 8)

    for grupo in grupos:
        doc.add_page_break()
        adicionar_cabecalho_grupo(doc, grupo)
        add_spacer(doc, 10, 6)

        for idx, item in enumerate(grupo.itens):
            adicionar_card_item(doc, item, is_zebra=(idx % 2 == 1))
            add_spacer(doc, 6, 2)

    add_spacer(doc, 12, 6)
    add_horizontal_rule(doc, Cor.VERDE_ESCURO, espaco_antes=4, espaco_depois=6)
    p_rodape = doc.add_paragraph()
    p_rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rodape.paragraph_format.space_before = Pt(0)
    p_rodape.paragraph_format.space_after = Pt(0)
    r_rod = p_rodape.add_run(
        "Gerado automaticamente — Pipeline RAG Local (Ollama + LangChain + ChromaDB)"
    )
    r_rod.italic = True; r_rod.font.size = Pt(9)
    set_run_color(r_rod, Cor.TEXTO_SUAVE)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()